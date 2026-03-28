from __future__ import annotations
from typing import List, Dict, Any, Optional, Tuple
from pathlib import Path
import os

BASE_DIR = Path(__file__).resolve().parents[2]
KNOWLEDGE_DIR = BASE_DIR / "knowledge"

DB_FILE_PATHS = [
    BASE_DIR / "Hackathon_Data sheet (1).xlsx",
    BASE_DIR / "List-of-Uploading-Documents-in-the-Online-Application-System.pdf",
    BASE_DIR / "the Major Catalog for Undergraduate International Students.docx",
]

DB_EXCLUDE_NAMES = {p.name for p in DB_FILE_PATHS}

DB_DOCS: List[Dict[str, Any]] = []


def _read_pdf_bytes(data: bytes) -> str:
    text = ""
    try:
        from pypdf import PdfReader
        import io
        reader = PdfReader(io.BytesIO(data))
        for page in reader.pages:
            t = page.extract_text() or ""
            if t:
                text += t
    except Exception:
        pass
    if not text.strip():
        try:
            from pdfminer.high_level import extract_text
            import io
            text = extract_text(io.BytesIO(data)) or ""
        except Exception:
            return ""
    return text


def _read_pdf(path: Path) -> str:
    try:
        return _read_pdf_bytes(path.read_bytes())
    except Exception:
        return ""


def _read_docx(path: Path) -> str:
    try:
        from docx import Document
        doc = Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs)
    except Exception:
        try:
            data = path.read_bytes()
            return data.decode("utf-8", errors="ignore")
        except Exception:
            return ""


def _read_xlsx(path: Path) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(filename=str(path), data_only=True, read_only=True)
        parts: List[str] = []
        for ws in wb.worksheets:
            parts.append(f"Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                vals = [str(v) if v is not None else "" for v in row]
                line = " | ".join(vals).strip()
                if line:
                    parts.append(line)
        return "\n".join(parts)
    except Exception:
        try:
            data = path.read_bytes()
            return data.decode("utf-8", errors="ignore")
        except Exception:
            return ""


def load_db() -> List[Dict[str, Any]]:
    global DB_DOCS
    docs: List[Dict[str, Any]] = []
    # static files
    for p in DB_FILE_PATHS:
        try:
            if not p.exists():
                continue
            name = p.name
            lower = name.lower()
            if lower.endswith(".pdf"):
                content = _read_pdf(p)
            elif lower.endswith(".docx"):
                content = _read_docx(p)
            elif lower.endswith(".xlsx") or lower.endswith(".xlsm"):
                content = _read_xlsx(p)
            else:
                try:
                    content = p.read_text(encoding="utf-8", errors="ignore")
                except Exception:
                    content = ""
            if content and content.strip():
                docs.append({"source": name, "path": str(p), "content": content})
        except Exception:
            continue
    # knowledge dir
    try:
        if KNOWLEDGE_DIR.exists():
            for p in KNOWLEDGE_DIR.iterdir():
                if not p.is_file():
                    continue
                name = p.name
                lower = name.lower()
                try:
                    if lower.endswith(".pdf"):
                        content = _read_pdf(p)
                    elif lower.endswith(".docx"):
                        content = _read_docx(p)
                    elif lower.endswith(".xlsx") or lower.endswith(".xlsm"):
                        content = _read_xlsx(p)
                    elif lower.endswith(".txt"):
                        content = p.read_text(encoding="utf-8", errors="ignore")
                    else:
                        continue
                    if content and content.strip():
                        docs.append({"source": name, "path": str(p), "content": content})
                except Exception:
                    continue
    except Exception:
        pass
    DB_DOCS = docs
    return DB_DOCS


def get_db_docs() -> List[Dict[str, Any]]:
    global DB_DOCS
    if not DB_DOCS:
        load_db()
    return DB_DOCS


def chunk_text(text: str, chunk_size: int = 1200, overlap: int = 200) -> List[str]:
    chunks: List[str] = []
    i = 0
    n = len(text)
    while i < n:
        j = min(i + chunk_size, n)
        chunks.append(text[i:j])
        if j == n:
            break
        i = max(j - overlap, i + 1)
    return chunks


def select_relevant_chunks(profile: Dict[str, Any], limit_chars: int = 8000) -> List[Dict[str, Any]]:
    # attempt embedding-based selection via Minimax embeddings if available
    minimax_key = os.getenv("MINIMAX_API_KEY")
    embed_model = os.getenv("MINIMAX_EMBED_MODEL") or "embedding-01"
    if minimax_key and embed_model:
        try:
            import httpx
            docs = get_db_docs()
            # build candidate chunks
            all_chunks: List[Tuple[str, int, str]] = []  # (source, idx, text)
            for d in docs:
                for idx, ch in enumerate(chunk_text(d["content"])):
                    all_chunks.append((d["source"], idx, ch))
            if not all_chunks:
                raise RuntimeError("no chunks")
            # create inputs: first the profile text, then chunks
            profile_text = []
            for k in ["gpa", "toefl", "ielts", "hsk"]:
                v = profile.get(k)
                if v is not None and str(v).strip():
                    profile_text.append(f"{k}:{v}")
            for k in ["majors", "cities", "skills"]:
                vs = profile.get(k) or []
                for v in vs:
                    if v:
                        profile_text.append(str(v))
            q_text = " | ".join(profile_text) or "international students requirements admissions documents"
            inputs = [q_text] + [t for (_, _, t) in all_chunks]
            base = os.getenv("MINIMAX_BASE_URL", "https://api.minimax.chat/v1").rstrip("/")
            headers = {"Authorization": f"Bearer {minimax_key}", "Content-Type": "application/json"}
            resp = httpx.post(f"{base}/text/embeddings", json={"model": embed_model, "input": inputs}, headers=headers, timeout=60)
            data = resp.json()
            vectors = []
            # expect data like {"data":[{"embedding":[...]}...]}
            if isinstance(data, dict) and "data" in data:
                for item in data["data"]:
                    vec = item.get("embedding")
                    if isinstance(vec, list):
                        vectors.append(vec)
            if len(vectors) != len(inputs):
                raise RuntimeError("embedding size mismatch")
            import math
            def dot(a, b):
                return sum(x*y for x, y in zip(a, b))
            def norm(a):
                return math.sqrt(sum(x*x for x in a)) or 1.0
            q_vec = vectors[0]
            scored = []
            for i, (src, idx, txt) in enumerate(all_chunks, start=1):
                v = vectors[i]
                s = dot(q_vec, v) / (norm(q_vec) * norm(v))
                scored.append((s, {"source": src, "index": idx, "text": txt}))
            # topic buckets
            buckets = {"documents": [], "majors": [], "policies": [], "other": []}
            for s, it in scored:
                src = (it["source"] or "").lower()
                if "upload" in src or "document" in src:
                    buckets["documents"].append((s, it))
                elif "catalog" in src or "major" in src:
                    buckets["majors"].append((s, it))
                elif "data" in src or "policy" in src or "sheet" in src:
                    buckets["policies"].append((s, it))
                else:
                    buckets["other"].append((s, it))
            for k in buckets:
                buckets[k].sort(key=lambda x: x[0], reverse=True)
            order = ["documents", "majors", "policies", "other"]
            picked: List[Dict[str, Any]] = []
            total = 0
            idx_map = {k: 0 for k in order}
            while total < limit_chars:
                advanced = False
                for k in order:
                    arr = buckets[k]
                    i = idx_map[k]
                    if i < len(arr):
                        item = arr[i][1]
                        L = len(item["text"])
                        if total + L > limit_chars:
                            continue
                        picked.append(item)
                        total += L
                        idx_map[k] = i + 1
                        advanced = True
                if not advanced:
                    break
            if picked:
                return picked
        except Exception:
            pass
    # fallback to keyword heuristic
    terms: List[str] = []
    for k in ["gpa", "toefl", "ielts", "hsk"]:
        v = profile.get(k)
        if v is not None and str(v).strip():
            terms.append(str(v).strip().lower())
    for k in ["majors", "cities", "skills"]:
        vs = profile.get(k) or []
        for v in vs:
            if v:
                terms.append(str(v).strip().lower())
    if not terms:
        terms = ["undergraduate", "master", "document", "catalog", "policy", "requirements", "international"]

    docs = get_db_docs()
    scored_chunks: List[Tuple[int, Dict[str, Any]]] = []
    for d in docs:
        chunks = chunk_text(d["content"])
        for idx, ch in enumerate(chunks):
            lower = ch.lower()
            score = sum(1 for t in terms if t in lower)
            if score > 0:
                scored_chunks.append((score, {"source": d["source"], "index": idx, "text": ch}))
    scored_chunks.sort(key=lambda x: x[0], reverse=True)
    buckets = {"documents": [], "majors": [], "policies": [], "other": []}
    for s, it in scored_chunks:
        src = (it["source"] or "").lower()
        if "upload" in src or "document" in src:
            buckets["documents"].append((s, it))
        elif "catalog" in src or "major" in src:
            buckets["majors"].append((s, it))
        elif "data" in src or "policy" in src or "sheet" in src:
            buckets["policies"].append((s, it))
        else:
            buckets["other"].append((s, it))
    for k in buckets:
        buckets[k].sort(key=lambda x: x[0], reverse=True)
    order = ["documents", "majors", "policies", "other"]
    picked: List[Dict[str, Any]] = []
    total = 0
    idx_map = {k: 0 for k in order}
    while total < limit_chars:
        advanced = False
        for k in order:
            arr = buckets[k]
            i = idx_map[k]
            if i < len(arr):
                item = arr[i][1]
                L = len(item["text"])
                if total + L > limit_chars:
                    continue
                picked.append(item)
                total += L
                idx_map[k] = i + 1
                advanced = True
        if not advanced:
            break
    if not picked:
        # fallback: take first chunk of each doc up to limit
        for d in docs:
            chunks = chunk_text(d["content"])
            if chunks:
                t = chunks[0]
                if total + len(t) > limit_chars:
                    break
                picked.append({"source": d["source"], "index": 0, "text": t})
                total += len(t)
            if total >= limit_chars:
                break
    return picked


def list_db_docs() -> List[Dict[str, Any]]:
    docs = []
    # static
    for p in DB_FILE_PATHS:
        try:
            if p.exists():
                st = p.stat()
                docs.append({"name": p.name, "path": str(p), "size": st.st_size, "location": "root"})
        except Exception:
            docs.append({"name": p.name, "path": str(p), "size": None, "location": "root"})
    # knowledge dir
    if KNOWLEDGE_DIR.exists():
        for p in KNOWLEDGE_DIR.iterdir():
            if p.is_file():
                try:
                    st = p.stat()
                    docs.append({"name": p.name, "path": str(p), "size": st.st_size, "location": "knowledge"})
                except Exception:
                    docs.append({"name": p.name, "path": str(p), "size": None, "location": "knowledge"})
    return docs
